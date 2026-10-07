import io
import json
import openpyxl
from nanopore.workflow import workflow_excel,read_settings_metadata
from nanopore.dal2 import REFERENCE,ANNOTATIONS


def workbook(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def text(book):
    return '\n'.join(str(cell.value) for sheet in book for row in sheet for cell in row if cell.value is not None)


def test_saved_read_settings_without_sample_results():
    metadata=dict(settings=dict(mode='Amplicon discovery',min_quality=11,min_length=200,max_length=5000,edits=2,anchor='ACGTACGT',gene='TGCATGCA',cluster_identity=.95),
        files={'PRIVATE_SAMPLE':[dict(name='private_reads.fastq.gz')]},tools={'vsearch':'2.30'},workflow_runtime={'pandas':'2.3'},workflow_completed_at='2026-10-07T12:00:00+02:00')
    book=workbook(workflow_excel(metadata,settings_source='Completed snapshot'))
    assert set(book.sheetnames)=={'About','Workflow','Settings','Percentages','Tools'}
    values={row[1]:row[2] for row in book['Settings'].iter_rows(min_row=2,values_only=True)}
    assert values['Minimum read Q score']==11
    assert values['VSEARCH clustering identity (fraction)']==.95
    assert 'PRIVATE_SAMPLE' not in text(book) and 'private_reads' not in text(book)
    assert 'not used in this discovery run' in text(book)


def test_alignment_workflow_does_not_invent_missing_fastq_settings():
    alignment=dict(reference=REFERENCE,annotations=ANNOTATIONS,trim=True,identity=.75,span=90,context_identity=.85)
    book=workbook(workflow_excel(alignment=alignment))
    assert len(list(book['Annotations'].iter_rows(min_row=2)))==9
    assert book['Reference']['B2'].value==REFERENCE
    assert 'Not recorded / not applicable' in text(book)
    assert 'minimum_downstream_context_identity' in text(book)
    assert '≥ max(85%' in text(book)


def test_imported_settings_whitelist_and_formula_safety():
    metadata=read_settings_metadata(json.dumps(dict(settings=dict(mode='Amplicon discovery',anchor='=1+1',min_quality=7,secret='do_not_export'),files={'sample':'no_export'})).encode())
    assert 'files' not in metadata and 'secret' not in metadata['settings']
    book=workbook(workflow_excel(metadata))
    cells=[cell for sheet in book for row in sheet for cell in row if cell.value=='=1+1']
    assert len(cells)==1 and cells[0].data_type=='s'


def test_workflow_available_with_complete_alignment_snapshot():
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    from test_summary import table
    from nanopore.summary import summarize_table
    frame=table();summary,members=summarize_table(frame,REFERENCE,ANNOTATIONS)
    page=Path(__file__).resolve().parents[1]/'pages'/'1_DAL2_Alignment_Viewer.py'
    app=AppTest.from_file(str(page))
    app.session_state['dal2_summary']=(summary,members,len(frame))
    app.session_state['dal2_dataset']=(frame,REFERENCE,ANNOTATIONS,True,.7,80,.8)
    app.run(timeout=30)
    assert not app.exception
    assert any('shared workflow' in x.label for x in app.expander)
