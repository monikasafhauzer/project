"""Sample-independent workflow workbook built from saved analysis snapshots."""
import io
import math
from datetime import datetime
from zoneinfo import ZoneInfo
from importlib.metadata import version, PackageNotFoundError
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

FASTQ_FIELDS = {
    'mode':'Analysis mode','min_quality':'Minimum read Q score','min_length':'Minimum input read length (bp)',
    'max_length':'Maximum input read length (bp)','retain_fasta':'Include FASTA with unknown quality',
    'anchor':'PCR forward/anchor primer (5′→3′)','gene':'PCR gene-specific reverse primer (5′→3′)',
    'edits':'Maximum substitutions + insertions + deletions per primer','keep_single':'Retain single-primer reads',
    'full':'Retain primer-containing amplicons','cluster_identity':'VSEARCH clustering identity (fraction)',
    'reference_identity':'Minimum minimap2 alignment identity (fraction)',
    'reference_coverage':'Minimum minimap2 reference coverage (fraction)',
    'full_coverage':'Full-length reference coverage threshold (fraction)'}


def runtime_info():
    info={}
    for package in ('streamlit','biopython','pandas','plotly','regex','openpyxl'):
        try: info[package]=version(package)
        except PackageNotFoundError: info[package]='Not available'
    return info


def read_settings_metadata(content):
    import json
    metadata=json.loads(content)
    if not isinstance(metadata,dict) or not isinstance(metadata.get('settings'),dict):
        raise ValueError('Use the analysis_settings.json exported by the read-analysis app.')
    settings=metadata['settings']
    if settings.get('mode') not in ('Amplicon discovery','Reference matching'):
        raise ValueError('The settings file has no recognized read-analysis mode.')
    # Only workflow fields; never copy files, paths, samples or result counts.
    result={'settings':{key:settings[key] for key in FASTQ_FIELDS if key in settings}}
    for field in ('tools','workflow_runtime','workflow_completed_at'):
        if field in metadata:
            if field in ('tools','workflow_runtime') and not isinstance(metadata[field],dict):
                raise ValueError(f'{field} must be a dictionary in the settings file.')
            result[field]=metadata[field]
    return result


def workflow_excel(read_metadata=None, alignment=None, settings_source='Not supplied'):
    read_metadata=read_metadata or {}
    settings=read_metadata.get('settings',{})
    mode=settings.get('mode')
    discovery=mode=='Amplicon discovery'
    matching=mode=='Reference matching'
    unknown=mode is None
    stages=[]
    def stage(number,name,rule,scope):
        stages.append(dict(step=number,stage=name,decision_rule=rule,applicability=scope))
    stage(1,'Basecalling / input','POD5 requires Dorado basecalling outside the app. Accept FASTQ/FASTA; gzip decompression preserves sequences.','Input preparation; basecalling is not performed by this app')
    stage(2,'Read quality','FASTQ Q = −10 log10(mean base error probability); require Q ≥ saved minimum. FASTA has no Q and follows the saved unknown-quality option.','Both read-analysis modes' if not unknown else 'General workflow; execution settings not supplied')
    stage(3,'Read length','Require saved minimum ≤ original read length ≤ saved maximum, before primer trimming. A read failing multiple checks is still one excluded read.','Both read-analysis modes')
    stage(4,'Primer detection and orientation','Search both strands. Primer edits include substitutions and indels per primer. Match anchor and reverse complement of the entered gene-specific oligo; prefer both-primer matches, lowest edits, then longest interior.','Discovery only'+(' — not used in this reference-matching run' if matching else ''))
    stage(5,'Read retention and extraction','Retain both-primer reads; optionally single-primer reads. Interior mode removes detected primer sites; single-primer mode keeps the corresponding flank. Full mode retains the primer-bounded amplicon, or whole oriented single-primer read.','Discovery only')
    stage(6,'Sequence clustering','VSEARCH --cluster_fast --iddef 1 --strand plus; sequence identity determines membership, never length alone. Samples are processed independently.','Discovery only')
    stage(7,'Preliminary consensus','VSEARCH --consout; one unpolished consensus per cluster. No polishing, chimera removal or biological isoform validation is performed.','Discovery only')
    stage(8,'Alternative reference matching','minimap2 -x map-ont -c --secondary=yes -p 0. Identity = matching bases / alignment block length; coverage = aligned reference span / reference length. Reads qualifying for multiple references are ambiguous and excluded from unique totals; full-length uses the saved coverage threshold.','Reference-matching mode only'+(' — not used in this discovery run' if discovery else ''))
    stage(9,'Grouped-table import','Analyze every imported CSV row or every nonempty Excel worksheet. Reject invalid rows/sheets or duplicate sample/amplicon IDs with an error, rather than silently skipping. Image row limits never limit full analysis.','DAL2 grouped-consensus workflow'+(' — saved alignment settings supplied' if alignment else ' — settings not supplied; not claimed executed'))
    stage(10,'Candidate TSO removal','Optionally remove a candidate full TSO, PCR-primer-plus-tail or GCAGAGTACATGGG tail near the oriented 5′ boundary, allowing one edit. Preserve removed sequence; never trim an internal motif. Terminal G ownership remains uncertain.','DAL2 alignment only')
    stage(11,'Consensus/reference alignment','Biopython local alignment, both orientations: match +2, mismatch −3, gap-open −5, gap-extension −1. Apply saved aligned-identity and reference-span thresholds; preserve unaligned ends.','DAL2 alignment only; identity is distinct from read clustering identity')
    stage(12,'uORF / start association','Find first intact annotated ATG with supported downstream reference context. Different 5′ prefixes do not imply earlier uORFs. Check 50-base context (≥40 covered bases), plus 12-base near-start context (≥10 covered bases, identity ≥ max(85%, saved context identity)). Require consistent candidate placement in detected alternate alignments. Strongly supported altered starts remain uncertain.','DAL2 alignment only')
    stage(13,'Grouped summary / display','Group per sample in reference uORF/start order. Include uncertain rows in totals. Image displays at most 100 selected rows; full reports retain every analyzed row. No sample-specific counts are included in this workbook.','Reporting')
    params=[]
    for key,label in FASTQ_FIELDS.items():
        params.append(dict(stage='Read analysis',setting=label,value=settings.get(key,'Not recorded / not applicable'),source=settings_source))
    if alignment:
        values=dict(trim_candidate_tso=alignment['trim'],minimum_aligned_identity=alignment['identity'],minimum_reference_span_bp=alignment['span'],minimum_downstream_context_identity=alignment['context_identity'],downstream_context_bases=50,minimum_covered_context_bases=40,near_start_context_bases=12,minimum_covered_near_start_bases=10,minimum_near_start_identity=max(.85,alignment['context_identity']),reference_length_bp=len(alignment['reference']))
        for key,value in values.items(): params.append(dict(stage='DAL2 alignment',setting=key,value=value,source='Completed full-table alignment snapshot'))
    else:
        params.append(dict(stage='DAL2 alignment',setting='Alignment settings',value='Not supplied; later alignment not claimed executed',source='Not recorded'))
    notes=[
        ('Purpose','One shared workflow and settings report; no sample names, filenames, individual reads, cluster counts or biological results.'),
        ('Execution evidence','This workbook documents saved settings and decision rules, not an audit of exclusions or proof that separate samples used identical settings.'),
        ('FASTQ settings source',settings_source),
        ('Settings relationship','Read settings supplied alongside imported consensus are not independently verified to originate from that file.' if alignment else 'Read settings come from the completed read-analysis snapshot when supplied.'),
        ('Snapshot','Uses completed-run settings, not current sidebar controls. Missing earlier-stage settings are explicitly not recorded.'),
        ('Generated at (Europe/Stockholm)',datetime.now(ZoneInfo('Europe/Stockholm')).isoformat()),
        ('Read-analysis completion time',read_metadata.get('workflow_completed_at','Not recorded')),
        ('Scientific meaning','PCR-derived sequencing-read abundance is not original RNA molecule abundance. ATG retention does not establish translation, complete ORFs or transcription start sites.'),
        ('Upstream sequence','Different/unaligned prefixes are preserved. Their origin is not assumed random; reference coordinates are not invented for them.'),
        ('Consensus / artifacts','Consensus is preliminary and unpolished. No PCR-chimera detection, artifact exclusion or biological isoform classification is performed.')]
    denominators=[
        ('Discovery: percent_quality_filtered','Cluster supporting reads / all quality-filtered reads × 100.'),
        ('Discovery: percent_assigned','Cluster supporting reads / all reads assigned to amplicon clusters × 100.'),
        ('Reference: percent_quality_filtered','Uniquely reference-matched reads / all quality-filtered reads × 100.'),
        ('Reference: percent_reference_assigned','Uniquely reference-matched reads / all uniquely reference-assigned reads × 100; ambiguous reads excluded.'),
        ('Grouped: percent_of_imported_supporting_reads','Group supporting reads / all imported supporting reads for the sample × 100; includes uncertain groups.'),
        ('Grouped: original percentages','Sum input percentages; preserve their original denominators. A subset input remains a subset; missing values are not zero.')]
    versions=[dict(tool=k,version=v,recorded_at='Read analysis completion') for k,v in read_metadata.get('tools',{}).items()]
    versions += [dict(tool=k,version=v,recorded_at='Read analysis completion') for k,v in read_metadata.get('workflow_runtime',{}).items()]
    versions += [dict(tool=k,version=v,recorded_at='Workbook export environment; not evidence of earlier run version') for k,v in runtime_info().items()]
    frames={'About':pd.DataFrame(notes,columns=['item','description']),'Workflow':pd.DataFrame(stages),'Settings':pd.DataFrame(params),'Percentages':pd.DataFrame(denominators,columns=['metric','denominator_rule']),'Tools':pd.DataFrame(versions)}
    if alignment:
        frames['Reference']=pd.DataFrame([dict(name='Annotated alignment reference',sequence=alignment['reference'],coordinates='1-based, inclusive; local to supplied reference')])
        frames['Annotations']=pd.DataFrame(alignment['annotations'])
    buffer=io.BytesIO()
    with pd.ExcelWriter(buffer,engine='openpyxl') as writer:
        for name,frame in frames.items():
            frame.to_excel(writer,sheet_name=name,index=False)
            sheet=writer.sheets[name]
            sheet.freeze_panes='A2'
            sheet.auto_filter.ref=sheet.dimensions
            for cell in sheet[1]:
                cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24556B')
            for column in sheet.columns:
                header=str(column[0].value)
                width=85 if header in ('description','decision_rule','denominator_rule','value','sequence') else 55 if header in ('setting','applicability','source','recorded_at') else 32
                sheet.column_dimensions[column[0].column_letter].width=width
                for cell in column[1:]:
                    cell.alignment=Alignment(wrap_text=True,vertical='top')
                    # Literal user text must never become an Excel formula.
                    if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')):
                        cell.data_type='s'
            for row in sheet.iter_rows(min_row=2):
                lines=max(max(1,math.ceil(len(str(cell.value or ''))/max(1,sheet.column_dimensions[cell.column_letter].width-3))) for cell in row)
                sheet.row_dimensions[row[0].row].height=min(400,max(32,lines*16+8))
    return buffer.getvalue()
